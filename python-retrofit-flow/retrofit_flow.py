import pandas as pd
import openpyxl
import warnings
import os
import re
import shutil
import threading
import logging
from datetime import datetime

import customtkinter as ctk
from tkinter import messagebox, filedialog

warnings.filterwarnings('ignore', category=UserWarning, module='openpyxl')

# ==========================================
# 0. CONFIGURAÇÃO DE LOGS E PASTAS
# ==========================================
pasta_logs = "Logs"
pasta_saida = "MBPs-preenchidas"
pasta_arquivados = "Relatorios-Processados"

for pasta in [pasta_logs, pasta_saida, pasta_arquivados]:
    if not os.path.exists(pasta):
        os.makedirs(pasta)

data_atual = datetime.now().strftime("%Y-%m-%d")
mes_ano = datetime.now().strftime("%m-%Y")

logger = logging.getLogger()
logger.setLevel(logging.INFO)
if logger.hasHandlers():
    logger.handlers.clear()

formato_log = logging.Formatter('%(asctime)s - %(levelname)s - %(message)s', datefmt='%d/%m/%Y %H:%M:%S')

log_erro = logging.FileHandler(os.path.join(pasta_logs, f"erros_{data_atual}.log"), encoding='utf-8')
log_erro.setLevel(logging.WARNING)
log_erro.setFormatter(formato_log)
logger.addHandler(log_erro)

# VERSÃO SANITIZADA PARA PORTFÓLIO: nome de arquivo e valores unitários
# abaixo são genéricos/de exemplo. Ajuste para o template e a tabela de
# preços reais do seu ambiente.
ARQUIVO_MODELO = "MODELO - RELATORIO DE ATIVOS - RETROFIT.xlsx"

# Valores contábeis de exemplo — troque pelos valores reais da sua tabela
# de ativos antes de usar em produção.
VALOR_PADRAO_NOTEBOOK = 0.00
VALOR_PADRAO_CPU = 0.00
VALOR_PADRAO_MONITOR = 0.00


# ==========================================
# 1. DESIGN TOKENS (paleta, tipografia)
# ==========================================
class Cores:
    BG = "#F3F4F7"            # fundo da janela
    SURFACE = "#FFFFFF"       # cartão de etapas
    BORDER = "#E4E6EB"

    INK = "#1C1E21"           # texto principal
    MUTED = "#6B7280"         # texto secundário

    BRAND = "#FFB612"         # amarelo Smart Fit (ação primária)
    BRAND_HOVER = "#F0A800"
    BRAND_INK = "#1C1E21"

    SECUNDARIO = "#EDEEF1"    # botão secundário / desabilitado
    SECUNDARIO_HOVER = "#E1E3E8"

    SUCESSO = "#16A34A"
    SUCESSO_HOVER = "#128A3E"
    ERRO = "#DC2626"
    AVISO = "#D97706"
    INFO = "#2563EB"

    TERM_BG = "#15171B"       # console
    TERM_DEFAULT = "#B6C2CF"
    TERM_SUCESSO = "#4ADE80"
    TERM_AVISO = "#FBBF24"
    TERM_ERRO = "#F87171"
    TERM_MUTED = "#5B6472"
    TERM_ACCENT = "#FFC629"


FONTE = "Segoe UI"
MONO = "Consolas"

ctk.set_appearance_mode("light")


# ==========================================
# 2. APLICAÇÃO
# ==========================================
class MBPApp(ctk.CTk):
    def __init__(self):
        super().__init__()

        self.arquivos_selecionados = []
        self.processando = False

        self.title("Gerador de MBP - Smart Fit")
        self.configure(fg_color=Cores.BG)
        # Descomente e aponte para um .ico próprio, se tiver um:
        # self.iconbitmap("icone.ico")

        self._definir_fontes()
        self._montar_layout()

        largura, altura = 780, 720
        self.minsize(680, 620)
        self._centralizar(largura, altura)

    # ------------------------------------------------------------
    # Construção da interface
    # ------------------------------------------------------------
    def _definir_fontes(self):
        self.fonte_titulo = ctk.CTkFont(family=FONTE, size=22, weight="bold")
        self.fonte_subtitulo = ctk.CTkFont(family=FONTE, size=12)
        self.fonte_secao = ctk.CTkFont(family=FONTE, size=13, weight="bold")
        self.fonte_botao = ctk.CTkFont(family=FONTE, size=13, weight="bold")
        self.fonte_corpo = ctk.CTkFont(family=FONTE, size=12)
        self.fonte_pequena = ctk.CTkFont(family=FONTE, size=11)
        self.fonte_mono = ctk.CTkFont(family=MONO, size=12)
        self.fonte_mono_pequena = ctk.CTkFont(family=MONO, size=11)

    def _centralizar(self, largura, altura):
        self.update_idletasks()
        sw = self.winfo_screenwidth()
        sh = self.winfo_screenheight()
        x = (sw - largura) // 2
        y = max((sh - altura) // 2 - 20, 0)
        self.geometry(f"{largura}x{altura}+{x}+{y}")

    def _montar_layout(self):
        self._montar_cabecalho()
        self._montar_card_etapas()
        self._montar_progresso()
        self._montar_console()
        self._montar_rodape()

    def _montar_cabecalho(self):
        frame = ctk.CTkFrame(self, fg_color="transparent")
        frame.pack(fill="x", padx=32, pady=(26, 6))

        ctk.CTkLabel(
            frame, text="Automação de MBP - Retrofit DELL",
            font=self.fonte_titulo, text_color=Cores.INK
        ).pack(anchor="w")

        ctk.CTkLabel(
            frame, text="Preenche as planilhas MBP automaticamente a partir dos checklists do retrofit Dell.",
            font=self.fonte_subtitulo, text_color=Cores.MUTED
        ).pack(anchor="w", pady=(2, 0))

    def _montar_card_etapas(self):
        card = ctk.CTkFrame(
            self, fg_color=Cores.SURFACE, corner_radius=14,
            border_width=1, border_color=Cores.BORDER
        )
        card.pack(fill="x", padx=32, pady=(14, 12))

        linha_botoes = ctk.CTkFrame(card, fg_color="transparent")
        linha_botoes.pack(fill="x", padx=20, pady=(20, 14))
        linha_botoes.grid_columnconfigure(0, weight=1)
        linha_botoes.grid_columnconfigure(1, weight=1)

        self.btn_selecionar = ctk.CTkButton(
            linha_botoes, text="1   Selecionar relatórios (Excel)",
            font=self.fonte_botao, height=46, corner_radius=9,
            fg_color=Cores.SECUNDARIO, hover_color=Cores.SECUNDARIO_HOVER,
            text_color=Cores.INK, command=self.selecionar_arquivos
        )
        self.btn_selecionar.grid(row=0, column=0, sticky="ew", padx=(0, 8))

        self.btn_processar = ctk.CTkButton(
            linha_botoes, text="2   Processar MBPs  ▶",
            font=self.fonte_botao, height=46, corner_radius=9,
            fg_color=Cores.SECUNDARIO, hover_color=Cores.SECUNDARIO_HOVER,
            text_color=Cores.MUTED, text_color_disabled=Cores.MUTED,
            state="disabled", command=self.iniciar_processamento
        )
        self.btn_processar.grid(row=0, column=1, sticky="ew", padx=(8, 0))

        linha_status = ctk.CTkFrame(card, fg_color="transparent")
        linha_status.pack(fill="x", padx=20, pady=(0, 4))

        self.dot_status = ctk.CTkLabel(
            linha_status, text="●", font=self.fonte_corpo,
            text_color=Cores.MUTED, width=14
        )
        self.dot_status.pack(side="left")

        self.lbl_status = ctk.CTkLabel(
            linha_status, text="Nenhum arquivo selecionado",
            font=self.fonte_corpo, text_color=Cores.MUTED
        )
        self.lbl_status.pack(side="left", padx=(2, 0))

        # Pré-visualização dos arquivos selecionados (some quando vazio)
        self.frame_lista_arquivos = ctk.CTkFrame(card, fg_color=Cores.BG, corner_radius=8)
        self.lista_arquivos = ctk.CTkTextbox(
            self.frame_lista_arquivos, height=64, fg_color=Cores.BG,
            text_color=Cores.MUTED, font=self.fonte_mono_pequena,
            corner_radius=8, wrap="none", activate_scrollbars=True
        )
        self.lista_arquivos.pack(fill="both", expand=True, padx=4, pady=4)
        self.lista_arquivos.configure(state="disabled")
        self._pack_opts_lista = {"fill": "x", "padx": 20, "pady": (0, 18)}

    def _montar_progresso(self):
        frame = ctk.CTkFrame(self, fg_color="transparent")
        frame.pack(fill="x", padx=32, pady=(0, 12))

        ctk.CTkLabel(
            frame, text="Progresso", font=self.fonte_secao, text_color=Cores.INK
        ).pack(anchor="w")

        self.progress = ctk.CTkProgressBar(
            frame, height=10, corner_radius=5,
            fg_color=Cores.SECUNDARIO, progress_color=Cores.SECUNDARIO
        )
        self.progress.set(0)
        self.progress.pack(fill="x", pady=(8, 4))

        self.lbl_progresso_texto = ctk.CTkLabel(
            frame, text="Aguardando início.", font=self.fonte_pequena, text_color=Cores.MUTED
        )
        self.lbl_progresso_texto.pack(anchor="w")

    def _montar_console(self):
        painel = ctk.CTkFrame(self, fg_color=Cores.TERM_BG, corner_radius=14)
        painel.pack(fill="both", expand=True, padx=32, pady=(0, 12))

        cabecalho = ctk.CTkFrame(painel, fg_color="transparent")
        cabecalho.pack(fill="x", padx=16, pady=(12, 4))

        ctk.CTkLabel(
            cabecalho, text="Console de execução", font=self.fonte_pequena,
            text_color=Cores.TERM_DEFAULT
        ).pack(side="left")

        btn_limpar = ctk.CTkButton(
            cabecalho, text="Limpar", width=60, height=22, corner_radius=6,
            fg_color="transparent", hover_color="#23262D", border_width=1,
            border_color=Cores.TERM_MUTED, text_color=Cores.TERM_MUTED,
            font=self.fonte_pequena, command=self._limpar_console
        )
        btn_limpar.pack(side="right")

        self.console = ctk.CTkTextbox(
            painel, fg_color=Cores.TERM_BG, text_color=Cores.TERM_DEFAULT,
            font=self.fonte_mono, corner_radius=10, wrap="word",
            activate_scrollbars=True
        )
        self.console.pack(fill="both", expand=True, padx=10, pady=(0, 10))

        for tag, cor in [
            ("default", Cores.TERM_DEFAULT), ("sucesso", Cores.TERM_SUCESSO),
            ("aviso", Cores.TERM_AVISO), ("erro", Cores.TERM_ERRO),
            ("muted", Cores.TERM_MUTED), ("accent", Cores.TERM_ACCENT),
        ]:
            self.console.tag_config(tag, foreground=cor)

        self.console.configure(state="disabled")
        self.log("Bem-vindo! Siga os passos acima para gerar suas planilhas.\n")

    def _montar_rodape(self):
        frame = ctk.CTkFrame(self, fg_color="transparent")
        frame.pack(fill="x", padx=32, pady=(0, 24))

        self.btn_abrir = ctk.CTkButton(
            frame, text="📂  Abrir pasta com planilhas prontas",
            font=self.fonte_botao, height=42, corner_radius=9,
            fg_color=Cores.SUCESSO, hover_color=Cores.SUCESSO_HOVER,
            text_color="#FFFFFF", command=self.abrir_pasta_saida
        )
        self.btn_abrir.pack(fill="x")

        ctk.CTkLabel(
            frame, text=f"As planilhas concluídas ficam salvas em: {pasta_saida}/",
            font=self.fonte_pequena, text_color=Cores.MUTED
        ).pack(anchor="w", pady=(6, 0))

    # ------------------------------------------------------------
    # Utilitários de UI (thread-safe: sempre agendados via .after)
    # ------------------------------------------------------------
    def log(self, mensagem, tipo=None):
        self.after(0, self._log_thread_principal, mensagem, tipo)

    def _log_thread_principal(self, mensagem, tipo):
        tag = tipo or self._detectar_tipo(mensagem)
        self.console.configure(state="normal")
        self.console.insert("end", mensagem + "\n", tag)
        self.console.see("end")
        self.console.configure(state="disabled")

    @staticmethod
    def _detectar_tipo(mensagem):
        texto = mensagem.strip()
        if not texto:
            return "default"
        if set(texto) <= set("=-—_ "):
            return "accent"
        maiusculo = texto.upper()
        if "⚠️" in texto or "ERRO" in maiusculo:
            return "erro" if "ERRO" in maiusculo else "aviso"
        if "✅" in texto or "🎉" in texto:
            return "sucesso"
        if "⏭️" in texto:
            return "muted"
        letras = [c for c in texto if c.isalpha()]
        if letras and all(c.isupper() for c in letras) and len(letras) > 6:
            return "accent"
        return "default"

    def _limpar_console(self):
        self.console.configure(state="normal")
        self.console.delete("1.0", "end")
        self.console.configure(state="disabled")

    def set_progress(self, atual, total):
        self.after(0, self._set_progress_thread_principal, atual, total)

    def _set_progress_thread_principal(self, atual, total):
        fracao = (atual / total) if total else 0
        self.progress.configure(progress_color=Cores.SECUNDARIO if fracao <= 0 else Cores.BRAND)
        self.progress.set(fracao)
        self.lbl_progresso_texto.configure(
            text=f"{atual} de {total} relatório(s) processado(s)."
        )

    def _resetar_progresso(self):
        self.progress.configure(progress_color=Cores.SECUNDARIO)
        self.progress.set(0)
        self.lbl_progresso_texto.configure(text="Aguardando início.")

    def set_status(self, estado, texto=None):
        mapa = {
            "idle": (Cores.MUTED, texto or "Nenhum arquivo selecionado"),
            "pronto": (Cores.INFO, texto or "Pronto para processar"),
            "processando": (Cores.AVISO, texto or "Processando..."),
            "concluido": (Cores.SUCESSO, texto or "Concluído"),
            "erro": (Cores.ERRO, texto or "Falha no processamento"),
        }
        cor, msg = mapa.get(estado, (Cores.MUTED, texto or ""))
        self.dot_status.configure(text_color=cor)
        self.lbl_status.configure(text=msg, text_color=Cores.MUTED if estado == "idle" else Cores.INK)

    def _atualizar_lista_arquivos(self):
        self.lista_arquivos.configure(state="normal")
        self.lista_arquivos.delete("1.0", "end")
        if self.arquivos_selecionados:
            for caminho in self.arquivos_selecionados:
                self.lista_arquivos.insert("end", f"📄 {os.path.basename(caminho)}\n")
            self.frame_lista_arquivos.pack(**self._pack_opts_lista)
        else:
            self.frame_lista_arquivos.pack_forget()
        self.lista_arquivos.configure(state="disabled")

    def _estilo_botao_processar(self, ativo):
        if ativo:
            self.btn_processar.configure(
                state="normal", fg_color=Cores.BRAND, hover_color=Cores.BRAND_HOVER,
                text_color=Cores.BRAND_INK
            )
        else:
            self.btn_processar.configure(
                state="disabled", fg_color=Cores.SECUNDARIO, hover_color=Cores.SECUNDARIO_HOVER,
                text_color=Cores.MUTED
            )

    def _estilo_botao_selecionar(self, ativo):
        self.btn_selecionar.configure(state="normal" if ativo else "disabled")

    # ------------------------------------------------------------
    # Ações do usuário
    # ------------------------------------------------------------
    def selecionar_arquivos(self):
        arquivos = filedialog.askopenfilenames(
            title="Selecione os relatórios baixados do Checklist Fácil",
            filetypes=[("Planilhas do Excel", "*.xlsx")]
        )
        if arquivos:
            self.arquivos_selecionados = list(arquivos)
            self._atualizar_lista_arquivos()
            self._estilo_botao_processar(ativo=True)
            self.set_status("pronto", f"{len(self.arquivos_selecionados)} arquivo(s) pronto(s) para processar")
            self.log(f"📂 {len(self.arquivos_selecionados)} relatório(s) carregado(s) e pronto(s) para processar.")

    def abrir_pasta_saida(self):
        if os.path.exists(pasta_saida):
            os.startfile(pasta_saida)  # Abre a pasta nativamente no Windows Explorer
        else:
            messagebox.showinfo("Aviso", "A pasta de saída ainda não existe ou está vazia.")

    def iniciar_processamento(self):
        if self.processando:
            return
        if not self.arquivos_selecionados:
            messagebox.showwarning("Aviso", "Por favor, clique em '1. Selecionar Relatórios' primeiro!")
            return

        self.processando = True
        self._estilo_botao_processar(ativo=False)
        self._estilo_botao_selecionar(ativo=False)
        self.set_status("processando")
        self._resetar_progresso()
        self.lbl_progresso_texto.configure(text="Iniciando...")

        thread = threading.Thread(target=self.processar_arquivos, daemon=True)
        thread.start()

    def _reabilitar_apos_erro_critico(self):
        def _fn():
            self.processando = False
            self._estilo_botao_selecionar(ativo=True)
            self._estilo_botao_processar(ativo=bool(self.arquivos_selecionados))
            self.set_status("erro")
        self.after(0, _fn)

    def _finalizar_processamento(self, msg_final):
        def _fn():
            self.processando = False
            self.arquivos_selecionados = []
            self._atualizar_lista_arquivos()
            self._estilo_botao_selecionar(ativo=True)
            self._estilo_botao_processar(ativo=False)
            self.set_status("concluido")
            messagebox.showinfo("Concluído", msg_final)
        self.after(0, _fn)

    # ------------------------------------------------------------
    # 3. LÓGICA DE NEGÓCIO — inalterada em relação à versão original,
    #    apenas adaptada para reportar progresso/estado à nova interface.
    # ------------------------------------------------------------
    def processar_arquivos(self):
        arquivos_selecionados = list(self.arquivos_selecionados)

        self.log("===================================================", "accent")
        self.log("    INICIANDO A AUTOMACAO DA SMART FIT...", "accent")
        self.log("===================================================\n", "accent")
        logging.info("--- INICIANDO PROCESSAMENTO VISUAL ---")

        arquivo_modelo = ARQUIVO_MODELO

        if not os.path.exists(arquivo_modelo):
            self.log(f"⚠️ ERRO: Arquivo modelo '{arquivo_modelo}' não encontrado.")
            messagebox.showerror("Erro", "Arquivo modelo não encontrado na pasta do programa!")
            self._reabilitar_apos_erro_critico()
            return

        arquivos_para_processar = []
        nomes_base_vistos = set()

        for arq in arquivos_selecionados:
            nome_arquivo = os.path.basename(arq)
            nome_limpo = re.sub(r'\s\(\d+\)', '', nome_arquivo)
            if nome_limpo not in nomes_base_vistos:
                nomes_base_vistos.add(nome_limpo)
                arquivos_para_processar.append(arq)
            else:
                self.log(f"⏭️ Arquivo repetido ignorado: {nome_arquivo}")

        self.log(f"\n🔎 Serão processados {len(arquivos_para_processar)} relatórios únicos.\n")
        unidades_processadas_total = 0
        total_arquivos = len(arquivos_para_processar)

        for indice, arquivo_checklist in enumerate(arquivos_para_processar, start=1):
            self.log(f"📄 Lendo relatório: {os.path.basename(arquivo_checklist)}")

            try:
                df = pd.read_excel(arquivo_checklist, header=1)
                todas_colunas = df.columns.tolist()

                if 'Unidade' not in df.columns:
                    self.log(f"   ⚠️ Coluna 'Unidade' não encontrada. Pulando arquivo.")
                    self.set_progress(indice, total_arquivos)
                    continue

                col_lote_patrimonio = None
                for c in todas_colunas:
                    if "Removido etiqueta de ativo fixo" in str(c):
                        col_lote_patrimonio = c
                        break

                colunas_equipamentos = []

                for i, col in enumerate(todas_colunas):
                    if "Service Tag" in str(col):
                        col_upper = str(col).upper()

                        if "NOTEBOOK" in col_upper:
                            tipo = "NOTEBOOK"
                            desc_padrao = "NOTEBOOK"
                            valor = VALOR_PADRAO_NOTEBOOK
                        elif "CPU" in col_upper:
                            tipo = "CPU"
                            desc_padrao = "CPU (OPTIPLEX)"
                            valor = VALOR_PADRAO_CPU
                        elif "MONITOR" in col_upper:
                            tipo = "MONITOR"
                            desc_padrao = "MONITOR"
                            valor = VALOR_PADRAO_MONITOR
                        else:
                            continue

                        col_pat = None
                        col_modelo = None
                        for j in range(1, 11):
                            if i + j < len(todas_colunas):
                                c_analisada = todas_colunas[i + j]
                                c_upper_an = str(c_analisada).upper()

                                if "ETIQUETA DE PATRIM" in c_upper_an and tipo in c_upper_an and not col_pat:
                                    col_pat = c_analisada

                                if "MODELO DO COMPUTADOR" in c_upper_an and tipo in ["CPU", "NOTEBOOK"] and not col_modelo:
                                    col_modelo = c_analisada

                        colunas_equipamentos.append({
                            "col_tag": col,
                            "col_pat": col_pat,
                            "col_modelo": col_modelo,
                            "desc_padrao": desc_padrao,
                            "valor": valor
                        })

                grupos_unidades = df.groupby('Unidade')

                for nome_unidade, grupo in grupos_unidades:
                    if pd.isna(nome_unidade) or str(nome_unidade).strip() == '':
                        continue

                    nome_unidade_str = str(nome_unidade).strip()
                    nome_seguro = re.sub(r'[\\/*?:"<>|]', "", nome_unidade_str)
                    arquivo_saida = os.path.join(pasta_saida, f"MBP - {nome_seguro} ({mes_ano}).xlsx")

                    partes_nome = nome_unidade_str.split("-", 1)
                    texto_sigla = partes_nome[0].strip() if len(partes_nome) > 0 else ""
                    texto_unidade = partes_nome[1].strip() if len(partes_nome) > 1 else ""

                    equipamentos_da_unidade = []

                    for index, row in grupo.iterrows():
                        fila_patrimonios = []
                        if col_lote_patrimonio and pd.notna(row[col_lote_patrimonio]):
                            fila_patrimonios.extend(re.findall(r'\d{4,}', str(row[col_lote_patrimonio])))

                        equip_temporario = []
                        modelo_geral_cpu = None

                        for equip in colunas_equipamentos:
                            val_tag = row[equip["col_tag"]]
                            val_tag_str = str(val_tag).strip().upper()

                            if val_tag_str not in ['N/A', 'NA', 'NAN', '', '0', '0.0', 'ILEGIVEL', 'ILEGÍVEL']:
                                tag_limpa = val_tag_str[:7]

                                pat_limpo = ""
                                if equip["col_pat"]:
                                    val_pat = row[equip["col_pat"]]
                                    if pd.notna(val_pat) and str(val_pat).strip().lower() not in ['n/a', 'na', '', 'ilegivel']:
                                        pat_limpo = str(val_pat).strip()
                                        if pat_limpo.endswith('.0'):
                                            pat_limpo = pat_limpo[:-2]

                                desc_final = equip["desc_padrao"]
                                if equip["col_modelo"]:
                                    val_mod = row[equip["col_modelo"]]
                                    if pd.notna(val_mod) and str(val_mod).strip().lower() not in ['n/a', 'na', '', 'ilegivel']:
                                        desc_final = str(val_mod).strip().upper()

                                if equip["desc_padrao"] == "CPU (OPTIPLEX)" and desc_final != "CPU (OPTIPLEX)":
                                    modelo_geral_cpu = desc_final

                                equip_temporario.append({
                                    "tag": tag_limpa,
                                    "patrimonio": pat_limpo,
                                    "desc": desc_final,
                                    "valor": equip["valor"],
                                    "desc_padrao": equip["desc_padrao"]
                                })

                        for eq in equip_temporario:
                            if eq["desc_padrao"] == "CPU (OPTIPLEX)" and modelo_geral_cpu:
                                eq["desc"] = modelo_geral_cpu

                            if not eq["patrimonio"] and fila_patrimonios:
                                eq["patrimonio"] = fila_patrimonios.pop(0)

                            if not eq["patrimonio"]:
                                eq["patrimonio"] = "0"

                        equipamentos_da_unidade.extend(equip_temporario)

                    if len(equipamentos_da_unidade) > 0:
                        workbook = openpyxl.load_workbook(arquivo_modelo)
                        planilha = workbook.active

                        def escrever_seguro(linha, coluna, texto, num_format=None):
                            cel_modificada = None
                            try:
                                cel_modificada = planilha.cell(row=linha, column=coluna)
                                cel_modificada.value = texto
                            except AttributeError:
                                for range_mesclado in planilha.merged_cells.ranges:
                                    if planilha.cell(row=linha, column=coluna).coordinate in range_mesclado:
                                        cel_modificada = planilha.cell(row=range_mesclado.min_row, column=range_mesclado.min_col)
                                        cel_modificada.value = texto
                                        break
                            if cel_modificada and num_format:
                                cel_modificada.number_format = num_format

                        escrever_seguro(11, 8, texto_sigla)
                        escrever_seguro(11, 12, texto_unidade)

                        linha_atual = 37
                        for equip in equipamentos_da_unidade:
                            escrever_seguro(linha_atual, 2, equip["patrimonio"])
                            escrever_seguro(linha_atual, 4, "DELL")
                            escrever_seguro(linha_atual, 5, equip["tag"])
                            escrever_seguro(linha_atual, 7, equip["desc"])
                            escrever_seguro(linha_atual, 11, equip["valor"], '"R$" #,##0.00')

                            linha_atual += 1

                        salvo = False
                        while not salvo:
                            try:
                                workbook.save(arquivo_saida)
                                salvo = True
                            except PermissionError:
                                resposta = messagebox.askretrycancel(
                                    "Arquivo Aberto",
                                    f"⛔ O arquivo '{os.path.basename(arquivo_saida)}' está aberto no Excel.\n\nFeche o arquivo e clique em 'Repetir'."
                                )
                                if not resposta:
                                    self.log(f"   ⚠️ Processamento cancelado.")
                                    break

                        if salvo:
                            self.log(f"   ✅ Gerado: {os.path.basename(arquivo_saida)}")
                            unidades_processadas_total += 1

                try:
                    shutil.copy2(arquivo_checklist, os.path.join(pasta_arquivados, os.path.basename(arquivo_checklist)))
                except Exception:
                    pass

            except Exception as e:
                self.log(f"   ⚠️ ERRO ao processar {os.path.basename(arquivo_checklist)}: {str(e)}")
                logging.error(f"Erro no arquivo {arquivo_checklist}", exc_info=True)

            self.set_progress(indice, total_arquivos)

        msg_final = f"Processo concluído! {unidades_processadas_total} planilhas geradas."
        self.log(f"\n🎉 {msg_final}", "sucesso")
        self.log("---------------------------------------------------\n", "accent")

        self._finalizar_processamento(msg_final)


if __name__ == "__main__":
    app = MBPApp()
    app.mainloop()